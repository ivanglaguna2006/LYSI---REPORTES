# -*- coding: utf-8 -*-
"""Gastos y mantenimiento: muestra en la aplicación los dos reportes de Excel que ya se hacen
(el mensual, con una hoja por mes, y el de resultados del año). No recalcula nada: lee cada hoja
tal como está y la pasa a página, respetando títulos, indicadores, tablas, colores, gráficos e imágenes.
Solo necesita openpyxl (corre igual en el navegador, dentro de Administración)."""
import os
import re
import json
import base64
import zipfile
import datetime
import posixpath
import html as H
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.utils import range_boundaries

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
GRIS = {"6B6B6B", "808080", "7F7F7F", "595959"}
SERIES = ["#2E7D32", "#9AA39C", "#C4841D", "#2A5C8A", "#5B4A8B", "#B33A3A"]


def e(s):
    return H.escape(str(s), quote=True).replace("\n", "<br>")


def _rgb(c):
    try:
        v = c.rgb
    except Exception:
        return None
    if isinstance(v, str) and len(v) >= 6:
        v = v[-6:].upper()
        return v
    return None


def _oscuro(hx):
    r, g, b = int(hx[0:2], 16), int(hx[2:4], 16), int(hx[4:6], 16)
    return (0.299 * r + 0.587 * g + 0.114 * b) < 150


def _num(v, dec, miles):
    s = f"{abs(v):,.{dec}f}"
    s = s.replace(",", "\x00").replace(".", ",").replace("\x00", "." if miles else "")
    return s


def fmt(v, nf):
    """Valor de una celda con el formato de número que tiene en el Excel (miles con punto, decimales con coma)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, datetime.datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, datetime.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    if isinstance(v, str):
        return v
    nf = nf or "General"
    if nf in ("General", "@"):
        if float(v).is_integer():
            return str(int(v))
        s = _num(v, 2, abs(v) >= 10000).rstrip("0").rstrip(",")
        return ("-" if v < 0 else "") + s
    secs = nf.split(";")
    neg = v < 0
    sec = secs[0]
    signo = "-" if neg else ""
    if neg and len(secs) > 1:
        sec, signo = secs[1], ""
    pre, suf, core, i, visto = "", "", "", 0, False
    while i < len(sec):
        ch = sec[i]
        if ch == "\\" and i + 1 < len(sec):
            lit = sec[i + 1]; i += 2
        elif ch == '"':
            j = sec.find('"', i + 1)
            j = len(sec) if j < 0 else j
            lit = sec[i + 1:j]; i = j + 1
        elif ch in "0#,.%?":
            core += ch; visto = True; i += 1
            continue
        elif ch in "_*":
            i += 2
            continue
        elif ch == "[":
            j = sec.find("]", i)
            i = (j + 1) if j >= 0 else len(sec)
            continue
        else:
            lit = ch; i += 1
        if visto:
            suf += lit
        else:
            pre += lit
    if not core:
        return signo + pre + suf
    pct = "%" in core
    x = float(v) * (100 if pct else 1)
    cuerpo = core.replace("%", "")
    dec = len(cuerpo.split(".")[1].replace(",", "")) if "." in cuerpo else 0
    miles = "," in cuerpo.split(".")[0]
    return signo + pre + _num(x, dec, miles) + ("%" if pct else "") + suf


class Celda(object):
    __slots__ = ("r", "c", "v", "cs", "fill", "color", "b", "i", "sz", "nf", "al", "link")


def leer_hoja(ws):
    """Devuelve {fila: [celdas]} con las celdas que tienen valor o color, ya resueltas las combinadas."""
    anc, cub = {}, set()
    for rg in ws.merged_cells.ranges:
        anc[(rg.min_row, rg.min_col)] = rg.max_col - rg.min_col + 1
        for r in range(rg.min_row, rg.max_row + 1):
            for c in range(rg.min_col, rg.max_col + 1):
                if (r, c) != (rg.min_row, rg.min_col):
                    cub.add((r, c))
    filas = {}
    for row in ws.iter_rows():
        for x in row:
            if (x.row, x.column) in cub:
                continue
            v = x.value
            if isinstance(v, str):
                v = v.strip() or None
            fill = None
            try:
                if x.fill is not None and x.fill.fill_type == "solid":
                    fill = _rgb(x.fill.fgColor)
                    if fill == "FFFFFF":
                        fill = None
            except Exception:
                fill = None
            if v is None and fill is None:
                continue
            k = Celda()
            k.r, k.c, k.v, k.cs, k.fill = x.row, x.column, v, anc.get((x.row, x.column), 1), fill
            f = x.font
            k.color = _rgb(f.color) if (f is not None and f.color is not None) else None
            if k.color == "000000":
                k.color = None
            k.b, k.i, k.sz = bool(f.b) if f is not None else False, bool(f.i) if f is not None else False, float(f.sz or 11) if f is not None else 11.0
            k.nf = x.number_format
            k.al = x.alignment.horizontal if x.alignment is not None else None
            k.link = None
            if x.hyperlink is not None and x.hyperlink.location:
                k.link = x.hyperlink.location
            elif x.hyperlink is not None:
                k.link = "externo"
            filas.setdefault(x.row, []).append(k)
    for r in filas:
        filas[r].sort(key=lambda k: k.c)
    return filas


# ───────────── gráficos e imágenes (se leen directo del archivo, sin depender de librerías de imagen) ─────────────
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}


def _rels(z, parte):
    ruta = posixpath.join(posixpath.dirname(parte), "_rels", posixpath.basename(parte) + ".rels")
    if ruta not in z.namelist():
        return {}
    out = {}
    for x in ET.fromstring(z.read(ruta)).findall("rel:Relationship", NS):
        t = x.get("Target")
        t = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(parte), t))
        out[x.get("Id")] = (x.get("Type", ""), t)
    return out


def dibujos(ruta):
    """{nombre de hoja: [(fila, tipo, dato)]} con tipo 'img' (bytes, extensión) o 'chart' (xml)."""
    out = {}
    try:
        z = zipfile.ZipFile(ruta)
        wrel = _rels(z, "xl/workbook.xml")
        for sh in ET.fromstring(z.read("xl/workbook.xml")).find("m:sheets", NS):
            rid = sh.get("{%s}id" % NS["r"])
            if rid not in wrel:
                continue
            parte = wrel[rid][1]
            for _, (tipo, dr) in _rels(z, parte).items():
                if not tipo.endswith("/drawing"):
                    continue
                drel = _rels(z, dr)
                for an in ET.fromstring(z.read(dr)):
                    de = an.find("xdr:from", NS)
                    if de is None:
                        continue
                    fila = int(de.find("xdr:row", NS).text) + 1
                    blip = an.find(".//a:blip", NS)
                    ch = an.find(".//c:chart", NS)
                    if blip is not None:
                        t = drel.get(blip.get("{%s}embed" % NS["r"]))
                        if t and t[1] in z.namelist():
                            out.setdefault(sh.get("name"), []).append((fila, "img", (z.read(t[1]), t[1].rsplit(".", 1)[-1].lower())))
                    elif ch is not None:
                        t = drel.get(ch.get("{%s}id" % NS["r"]))
                        if t and t[1] in z.namelist():
                            out.setdefault(sh.get("name"), []).append((fila, "chart", z.read(t[1])))
    except Exception:
        pass
    for k in out:
        out[k].sort(key=lambda x: x[0])
    return out


def _rango(wb, ref):
    """Valores de una referencia tipo 'Hoja'!$C$6:$C$14."""
    if not ref or "!" not in ref:
        return []
    hoja, rg = ref.rsplit("!", 1)
    hoja = hoja.strip("'").replace("''", "'")
    if hoja not in wb.sheetnames:
        return []
    c0, r0, c1, r1 = range_boundaries(rg.replace("$", ""))
    ws = wb[hoja]
    return [ws.cell(r, c).value for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)]


def _corto(v):
    a = abs(v)
    if a >= 1e6:
        return _num(v / 1e6, 0 if a >= 1e7 else 1, True) + " M"
    if a >= 1e4:
        return _num(v / 1e3, 0, True) + " mil"
    return _num(v, 0 if float(v).is_integer() else 1, True)


def chart_svg(wb, xml):
    """Gráfico del Excel redibujado como barras agrupadas con los mismos datos."""
    try:
        root = ET.fromstring(xml)
        tit = " ".join(t.text or "" for t in root.findall("c:chart/c:title//a:t", NS)).strip()
        series, cats = [], []
        for s in root.findall(".//c:ser", NS):
            nom = s.find("c:tx//c:f", NS)
            nomv = s.find("c:tx//c:v", NS)
            nombre = (_rango(wb, nom.text) or [""])[0] if nom is not None else (nomv.text if nomv is not None else "")
            cf = s.find("c:cat//c:f", NS)
            vf = s.find("c:val//c:f", NS)
            if cf is not None and not cats:
                cats = [fmt(x, None) for x in _rango(wb, cf.text)]
            vals = [x if isinstance(x, (int, float)) else 0 for x in _rango(wb, vf.text)] if vf is not None else []
            if vals:
                series.append((str(nombre or ""), vals))
        if not series:
            return ""
        n = max(len(v) for _, v in series)
        cats = (cats + [""] * n)[:n]
        W, Hh, ml, mb, mt = 980, 320, 64, 34, 14
        mx = max(max(v) for _, v in series) or 1
        paso = 10 ** (len(str(int(mx))) - 1)
        tope = (int(mx / paso) + 1) * paso
        if tope / paso > 6:
            paso *= 2
        aw = (W - ml - 10) / n
        bw = min(34, aw * 0.8 / len(series))
        o = [f'<svg class="chart" viewBox="0 0 {W} {Hh}" role="img" aria-label="{e(tit or "Gráfico")}">']
        y0 = Hh - mb
        k = 0
        while k * paso <= tope + 1e-9:
            y = y0 - (y0 - mt) * (k * paso) / tope
            o.append(f'<line x1="{ml}" x2="{W - 10}" y1="{y:.1f}" y2="{y:.1f}" stroke="#dfe4df"/><text class="ax" x="{ml - 6}" y="{y + 4:.1f}" text-anchor="end">{e(_corto(k * paso))}</text>')
            k += 1
        for i in range(n):
            cx = ml + aw * i + aw / 2
            for j, (_, v) in enumerate(series):
                val = v[i] if i < len(v) else 0
                h = (y0 - mt) * max(val, 0) / tope
                x = cx - bw * len(series) / 2 + bw * j
                o.append(f'<rect x="{x:.1f}" y="{y0 - h:.1f}" width="{bw - 2:.1f}" height="{h:.1f}" fill="{SERIES[j % len(SERIES)]}"><title>{e(cats[i])} · {e(series[j][0])}: {e(_num(val, 0, True))}</title></rect>')
            o.append(f'<text class="ax" x="{cx:.1f}" y="{y0 + 16}" text-anchor="middle">{e(str(cats[i])[:12].capitalize())}</text>')
        o.append("</svg>")
        ley = "".join(f'<span class="le"><i style="background:{SERIES[j % len(SERIES)]}"></i>{e(nm)}</span>' for j, (nm, _) in enumerate(series) if nm)
        return f'<figure class="gfig">{"<figcaption>" + e(tit) + "</figcaption>" if tit else ""}{"".join(o)}<div class="ley">{ley}</div></figure>'
    except Exception:
        return ""


# ───────────── de hoja a página ─────────────
def _vals(cs):
    return [k for k in cs if k.v is not None]


def _tipo(r, filas):
    cs = _vals(filas.get(r, []))
    if not cs:
        return "blanco"
    if r == 1:
        return "titulo"
    if len(cs) == 1 and cs[0].link and str(cs[0].v).startswith("←"):
        return "saltar"
    if all(k.sz >= 16 for k in cs):
        return "kpi"
    if len(cs) >= 2 and all(k.fill and k.b and k.color == "FFFFFF" and isinstance(k.v, str) for k in cs):
        return "th"
    if len(cs) == 1:
        k = cs[0]
        if k.fill and k.b and k.color == "FFFFFF":
            return "barra"
        if k.b and k.sz >= 13:
            return "h2"
        if k.cs >= 3 or (isinstance(k.v, str) and len(k.v) > 60) or k.c > 1:
            return "p"
        return "p1"          # una sola celda corta en la primera columna: puede ser fila de tabla
    return "fila"


class Estilos(object):
    """Colores de fila como clases, para no repetir el estilo en cada celda."""
    def __init__(self):
        self.m = {}

    def clase(self, hx):
        if hx not in self.m:
            self.m[hx] = "gf%d" % len(self.m)
        return self.m[hx]

    def css(self):
        o = []
        for hx, c in self.m.items():
            o.append(f".gst tbody tr.{c} td{{background:#{hx}{';color:#fff;font-weight:600' if _oscuro(hx) else ''}}}")
        return "".join(o)


def _celda_txt(k, ids):
    t = fmt(k.v, k.nf)
    h = e(t)
    if k.link and k.link != "externo":
        hoja = k.link.rsplit("!", 1)[0].strip("'")
        if hoja in ids:
            return f'<button type="button" class="lnk" data-ir="{ids[hoja]}">{h}</button>'
    if isinstance(k.v, str) and len(t) > 34:
        h = f'<span class="obs">{h}</span>'
    return h


def _td(k, ids, fila_fill, tag="td"):
    num = isinstance(k.v, (int, float)) and not isinstance(k.v, bool)
    cls, st = [], []
    if num or k.al == "right":
        cls.append("r")
    elif k.al == "center":
        cls.append("c")
    if k.fill and k.fill != fila_fill:
        st.append(f"background:#{k.fill}")
        if _oscuro(k.fill):
            st.append("color:#fff")
    if k.color and k.color != "FFFFFF":
        if k.color in GRIS:
            cls.append("gr")
        else:
            st.append(f"color:#{k.color}")
    if k.b and tag == "td":
        st.append("font-weight:600")
    a = (f' class="{" ".join(cls)}"' if cls else "") + (f' style="{";".join(st)}"' if st else "") + (f' colspan="{k.cs}"' if k.cs > 1 else "")
    return f"<{tag}{a}>{_celda_txt(k, ids)}</{tag}>"


def _tr(cs, c0, c1, ids, est, tag="td"):
    fills = {k.fill for k in cs}
    ff = list(fills)[0] if len(fills) == 1 and None not in fills and tag == "td" and _cubre(cs, c0, c1) else None
    o, c = [], c0
    for k in cs:
        if k.c < c:
            continue
        if k.c > c:
            o.append(f'<{tag} class="gap"{" colspan=%d" % (k.c - c) if k.c - c > 1 else ""}></{tag}>')
        o.append(_td(k, ids, ff, tag))
        c = k.c + k.cs
    if c <= c1:
        o.append(f'<{tag} class="gap"{" colspan=%d" % (c1 - c + 1) if c1 - c + 1 > 1 else ""}></{tag}>')
    return f'<tr{" class=" + est.clase(ff) if ff else ""}>{"".join(o)}</tr>'


def _cubre(cs, c0, c1):
    return sum(k.cs for k in cs) >= (c1 - c0 + 1) * 0.6


def _parrafos(ps):
    """ps: lista de celdas de texto seguidas."""
    o, lista = [], []

    def cerrar():
        if lista:
            o.append('<ul class="hechos">' + "".join(lista) + "</ul>")
            del lista[:]
    for k in ps:
        t = fmt(k.v, k.nf)
        if t.startswith("•"):
            lista.append(f"<li>{e(t[1:].strip())}</li>")
            continue
        cerrar()
        if t.startswith("⚠"):
            o.append(f'<p class="aviso">{e(t)}</p>')
        elif k.color in GRIS or k.i or k.sz < 11:
            o.append(f'<p class="nota">{e(t)}</p>')
        else:
            st = f' style="color:#{k.color}"' if k.color and k.color != "FFFFFF" else ""
            o.append(f'<p class="gp"{st}>{"<b>" if k.b else ""}{e(t)}{"</b>" if k.b else ""}</p>')
    cerrar()
    return "".join(o)


def hoja_html(wb, ws, ids, dib, est):
    """Devuelve (titulo, subtitulo, color, html del cuerpo)."""
    filas = leer_hoja(ws)
    ult = max(filas) if filas else 0
    tipo = {r: _tipo(r, filas) for r in range(1, ult + 3)}
    pend = list(dib)
    B = []            # bloques: (clase, html) ; clase en h2 | barra | tabla | otro
    titulo, sub, color = ws.title, "", None

    def soltar(hasta):
        while pend and pend[0][0] <= hasta:
            _, t, d = pend.pop(0)
            if t == "img":
                mime = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif"}.get(d[1])
                if mime:
                    B.append(("otro", f'<figure class="gfig"><img alt="Gráfico del reporte" loading="lazy" src="data:{mime};base64,{base64.b64encode(d[0]).decode()}"></figure>'))
            else:
                B.append(("otro", chart_svg(wb, d)))

    r = 1
    while r <= ult:
        t = tipo[r]
        cs = _vals(filas.get(r, []))
        if t in ("blanco", "saltar"):
            r += 1
            continue
        soltar(r)
        if t == "titulo":
            titulo, color = fmt(cs[0].v, cs[0].nf), cs[0].fill
            if tipo.get(2) in ("p", "p1"):
                sub = fmt(_vals(filas[2])[0].v, None)
                r = 3
            else:
                r = 2
            continue
        # indicadores: rótulos arriba, valor grande, aclaración abajo
        if t == "kpi" or (t in ("fila", "p", "p1") and tipo.get(r + 1) == "kpi" and all(k.sz <= 11 for k in cs)):
            rot = {}
            if t != "kpi":
                rot = {k.c: k for k in cs}
                r += 1
            valores = _vals(filas[r])
            notas = {}
            if tipo.get(r + 1) in ("fila", "p", "p1") and tipo.get(r + 2) != "kpi":
                nx = _vals(filas[r + 1])
                if all(k.sz <= 10 for k in nx) and {k.c for k in nx} <= {k.c for k in valores}:
                    notas = {k.c: k for k in nx}
            o = []
            for k in valores:
                st = f' style="color:#{k.color}"' if k.color and k.color not in GRIS else ""
                o.append(f'<div class="kpi"><div class="kl">{e(fmt(rot[k.c].v, None)) if k.c in rot else ""}</div><div class="kv"{st}>{e(fmt(k.v, k.nf))}</div>'
                         f'{"<div class=kn>" + e(fmt(notas[k.c].v, None)) + "</div>" if k.c in notas else ""}</div>')
            B.append(("otro", f'<div class="kpis gk">{"".join(o)}</div>'))
            r += 2 if notas else 1
            continue
        if t == "h2":
            k = cs[0]
            B.append(("h2", (fmt(k.v, None), k.color)))
            r += 1
            continue
        if t == "barra":
            k = cs[0]
            B.append(("barra", (fmt(k.v, None), k.fill, k.sz)))
            r += 1
            continue
        if t in ("p", "p1"):
            ps = []
            while r <= ult and tipo[r] in ("p", "p1") and tipo.get(r + 1) != "kpi":
                ps.append(_vals(filas[r])[0])
                r += 1
            if not ps:      # rótulo suelto antes de un indicador sin valores: se muestra como texto
                ps.append(cs[0]); r += 1
            B.append(("otro", _parrafos(ps)))
            continue
        # tabla: encabezados (uno o más renglones) y cuerpo hasta el primer renglón en blanco o título
        cab, cuerpo = [], []
        while r <= ult and tipo[r] == "th":
            cab.append(filas[r]); r += 1
        while r <= ult and tipo[r] in ("fila", "p1") and tipo.get(r + 1) != "kpi":
            cuerpo.append(filas[r]); r += 1
        if not cab and not cuerpo:      # fila que antecede a indicadores pero no es rótulo
            cuerpo.append(filas[r]); r += 1
        todas = cab + cuerpo
        c0 = min(k.c for f in todas for k in f)
        c1 = max(k.c + k.cs - 1 for f in todas for k in f)
        suelta = not cab and all(k.b and k.sz >= 13 for f in cuerpo for k in _vals(f))
        h = '<div class="tw sl"><table>' if suelta else '<div class="tw"><table>'
        if cab:
            h += "<thead>" + "".join(_tr(f, c0, c1, ids, est, "th") for f in cab) + "</thead>"
        h += "<tbody>" + "".join(_tr(f, c0, c1, ids, est) for f in cuerpo) + "</tbody></table></div>"
        B.append(("tabla", (h, len(cuerpo))))
    soltar(10 ** 9)

    # armado: secciones por h2; barras seguidas de tabla pasan a ser bloques plegables cuando son muchas
    pares = sum(1 for i in range(len(B) - 1) if B[i][0] == "barra" and B[i + 1][0] == "tabla")
    plegar = pares > 8
    o, abierta, i = [], False, 0
    while i < len(B):
        cl, d = B[i]
        if cl == "h2":
            if abierta:
                o.append("</section>")
            st = f' style="border-left-color:#{d[1]}"' if d[1] and d[1] not in GRIS else ""
            o.append(f"<section><h2{st}>{e(d[0])}</h2>")
            abierta = True
        elif cl == "barra":
            if plegar and i + 1 < len(B) and B[i + 1][0] == "tabla" and d[2] < 13:
                o.append(f'<details class="hist gd"><summary><span class="pt" style="background:#{d[1]}"></span>{e(d[0])}</summary>{B[i + 1][1][0]}</details>')
                i += 2
                continue
            o.append(f'<h3 class="barra{" b2" if d[2] < 13 else ""}" style="background:#{d[1]}">{e(d[0])}</h3>')
        elif cl == "tabla":
            o.append(d[0])
        else:
            o.append(d)
        i += 1
    if abierta:
        o.append("</section>")
    return titulo, sub, color, "".join(o)


def _clasificar(wb):
    hs = [h.lower() for h in wb.sheetnames]
    if sum(1 for h in hs if h in MESES) >= 1:
        return "mensual"
    if "resumen" in hs or "gasto mensual" in hs:
        return "resultados"
    return None


def control_mensual(wb):
    """Tabla de control: el resumen de cada mes tal como figura en la hoja Índice."""
    idx = [h for h in wb.sheetnames if h.lower() not in MESES]
    if not idx:
        return None
    filas = leer_hoja(wb[idx[0]])
    for r in sorted(filas):
        if _tipo(r, filas) == "th":
            cab = [fmt(k.v, None) for k in _vals(filas[r])]
            cols = [k.c for k in _vals(filas[r])]
            cuerpo = []
            r += 1
            while r in filas and _tipo(r, filas) in ("fila", "p1"):
                d = {k.c: fmt(k.v, k.nf) for k in _vals(filas[r])}
                cuerpo.append([d.get(c, "") for c in cols])
                r += 1
            return {"cab": cab, "filas": cuerpo}
    return None


def procesar(carpeta):
    """Lee los Excel de gastos que haya en la carpeta y arma el HTML del apartado. Devuelve un dict."""
    out = {"html": None, "errores": [], "archivos": {}, "control": None, "hojas": 0, "descripcion": ""}
    libros = {}
    for f in sorted(os.listdir(carpeta)) if os.path.isdir(carpeta) else []:
        if not f.lower().endswith((".xlsx", ".xlsm")) or f.startswith("~$"):
            continue
        ruta = os.path.join(carpeta, f)
        try:
            wb = openpyxl.load_workbook(ruta, data_only=True)
        except Exception as ex:
            out["errores"].append(f"No se pudo abrir {f}. Detalle: {type(ex).__name__}: {ex}")
            continue
        t = _clasificar(wb)
        if t is None:
            out["errores"].append(f"{f}: no reconozco este Excel como el reporte mensual (una hoja por mes) ni como el de resultados (hoja Resumen).")
            continue
        if t in libros:
            out["errores"].append(f"{f}: hay dos archivos del mismo tipo ({'mensual' if t == 'mensual' else 'resultados'}). Dejá uno solo.")
            continue
        libros[t] = (f, ruta, wb)
        out["archivos"][t] = f
    if not libros:
        return out

    # identificadores de cada hoja
    planes = []     # (tipo, hoja, id, etiqueta)
    if "resultados" in libros:
        for i, h in enumerate(libros["resultados"][2].sheetnames):
            planes.append(("resultados", h, "gr%02d" % i, h))
    if "mensual" in libros:
        for i, h in enumerate(libros["mensual"][2].sheetnames):
            planes.append(("mensual", h, "gm%02d" % i, h))
    est = Estilos()
    partes, vacias = [], 0
    for tipo in ("resultados", "mensual"):
        if tipo not in libros:
            continue
        f, ruta, wb = libros[tipo]
        ids = {h: i for t, h, i, _ in planes if t == tipo}
        dib = dibujos(ruta)
        chips = ""
        for t, h, i, et in planes:
            if t != tipo:
                continue
            ws = wb[h]
            titulo, sub, color, cuerpo = hoja_html(wb, ws, ids, dib.get(h, []), est)
            if not cuerpo:
                vacias += 1
            nav = ""
            if tipo == "resultados":
                nav = '<nav class="chips" aria-label="Hojas del reporte de resultados">' + "".join(
                    f'<button type="button" data-ir="{i2}"{" aria-current=true" if i2 == i else ""}>{e(et2)}</button>'
                    for t2, _, i2, et2 in planes if t2 == "resultados") + "</nav>"
            partes.append(f'<div class="apartado gst" id="{i}"{" data-grupo=gr" if tipo == "resultados" else ""} hidden>'
                          f'<header class="aph"><h1>{e(titulo)}</h1>{"<p>" + e(sub) + "</p>" if sub else ""}</header>{nav}{cuerpo}</div>')
            out["hojas"] += 1
    if vacias:
        out["errores"].append(f"{vacias} hoja(s) llegaron sin contenido: si el Excel tiene fórmulas, abrilo y guardalo en Excel antes de subirlo.")
    tabs = []
    if "resultados" in libros:
        tabs.append('<button type="button" data-ir="gr00" data-grupo="gr">Resultados del año</button>')
    if "mensual" in libros:
        if tabs:
            tabs.append('<span class="sep"></span>')
        for t, h, i, et in planes:
            if t == "mensual":
                tabs.append(f'<button type="button" data-ir="{i}">{e("Mes a mes" if h.lower() not in MESES else et)}</button>')
        out["control"] = control_mensual(libros["mensual"][2])
    inicio = "gr00" if "resultados" in libros else [i for t, h, i, _ in planes if t == "mensual"][-1]
    meses = [h for t, h, _, _ in planes if t == "mensual" and h.lower() in MESES]
    out["descripcion"] = (f"{meses[0]} a {meses[-1].lower()}" if len(meses) > 1 else (meses[0] if meses else "")) + (
        " · resultados del año" if "resultados" in libros else "")
    out["html"] = f'<template data-tabs>{"".join(tabs)}</template><div data-inicio="{inicio}"></div><style>{est.css()}</style>' + "".join(partes)
    return out


def procesar_json(carpeta):
    return json.dumps(procesar(carpeta), ensure_ascii=False)
