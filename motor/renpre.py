# -*- coding: utf-8 -*-
"""Control RENPRE: cruza los remitos de LYSI (lo que realmente se transportó) contra las declaraciones
presentadas ante el RENPRE, siguiendo la especificación «Control RENPRE — cruce remitos vs. declaraciones».

Corre igual en escritorio y dentro del navegador (Pyodide). No guarda nada: lee los dos archivos, cruza y
devuelve el resultado (listas y el HTML de la pantalla)."""
import re
import json
import html
import datetime
import unicodedata
from collections import defaultdict

VERSION = "1.0"

# ── Reglas (por ahora acá; la idea es que pasen a configuración editable desde la app) ──
REGLAS = {
    # Observaciones manda sobre el cliente: la carga sale de otro lado.
    "retira": [["retira unipar", "unipar"], ["retira inquimex", "inquimex"]],
    # Palabra clave en Cliente -> origen esperado. Se recorre en orden.
    "origen_por_cliente": [["unipar", "unipar"], ["ypf", "ypf"], ["sunchemical", "sun chemical"], ["thinner", "thinner tede"],
                           ["thiner", "thinner tede"], ["inquimex", "inquimex"], ["alper", "alper"], ["vasa", "vidrieria"],
                           ["houghton", "houghton"], ["carfi", "carfi"], ["brenntag", "brenntag"], ["balcarce", "balcarce"],
                           ["unilever", "unipar"], ["calsa", "calsa"]],
    # Substring en el nombre del producto -> familia. Se recorre en orden; sin match, primeros 12 caracteres.
    "familias": [["soda", "soda"], ["solvente", "solvente"], ["amoniac", "amoniacal"], ["tinta", "tinta"], ["xilen", "xileno"],
                 ["diluyente", "diluyente"], ["aguarras", "diluyente"], ["tolueno", "tolueno"], ["sebo", "sebo"],
                 ["exomax", "exomax"], ["lab", "lab"]],
    # Familias del remito para las que no se exige coincidencia (Thinner Tede mezcla patas de distinto solvente).
    "familias_libres": ["diluyente", "xileno"],
    "stopwords": ("de la el los las del y e sa s.a saic srl ltda coop dt t hs ton sale termino planta pto gral san int argentina "
                  "sociedad anonima comercial industrial financiera inmob saica cia sacif sau saci sacei caseros").split(),
    "aliases": {"vasa": "vidrieria", "diranza": "diransa", "coop": "cooperativa"},
    "tolerancia_kg": 300.0,
    "tolerancia_pct": 3.0,
}

MESES = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]


# ───────────────────────── utilidades ─────────────────────────
def _txt(v):
    if v is None:
        return ""
    try:
        import pandas as pd
        if pd.isna(v):
            return ""
    except Exception:
        pass
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _num(v):
    """Número desde celda: acepta 1.234,5 y 1234.5. None si no es número."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        try:
            import math
            if isinstance(v, float) and math.isnan(v):
                return None
        except Exception:
            pass
        return float(v)
    s = str(v).strip().replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _fecha(v):
    """Devuelve (iso 'YYYY-MM-DD' o '', texto para mostrar)."""
    if v is None or v == "":
        return "", ""
    if isinstance(v, datetime.datetime) or isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d"), v.strftime("%d/%m/%Y")
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return m.group(0), f"{m.group(3)}/{m.group(2)}/{m.group(1)}"
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})", s)
    if m:
        a = m.group(3) if len(m.group(3)) == 4 else "20" + m.group(3)
        return f"{a}-{int(m.group(2)):02d}-{int(m.group(1)):02d}", f"{int(m.group(1)):02d}/{int(m.group(2)):02d}/{a}"
    return "", s[:10]


def _sin_acentos(s):
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def tokens(texto, reglas=None):
    """3.8: minúsculas, puntuación a espacio, sin stopwords, sin números sueltos, sin tokens de 2 letras o menos, con aliases."""
    R = reglas or REGLAS
    t = _sin_acentos(str(texto or "").lower()).replace(".", " ").replace(",", " ")
    t = re.sub(r"[^a-z0-9 ]", " ", t)
    stop = set(R["stopwords"])
    out = set()
    for w in t.split():
        if w in stop or w.isdigit() or len(w) <= 2:
            continue
        out.add(R["aliases"].get(w, w))
    return out


def clave_doc(s):
    """Número de documento normalizado para cruzar: sin espacios, sin .0, sin ceros a la izquierda."""
    s = _txt(s).replace(" ", "")
    if re.match(r"^\d+\.0+$", s):
        s = s.split(".")[0]
    return s.lstrip("0") or ("0" if s else "")


def partes_remito(remito):
    """3.1: '53861-53862' -> ['53861', '53862']; '77528' -> ['77528']."""
    s = re.sub(r"\s+", "", _txt(remito))
    return [p for p in re.split(r"[-/]", s) if p]


def familia(nombre, reglas=None):
    R = reglas or REGLAS
    n = _sin_acentos(str(nombre or "").lower())
    for sub, fam in R["familias"]:
        if sub in n:
            return fam
    return n[:12].strip()


def origen_esperado(cliente, obs, reglas=None):
    """3.3: Observaciones ('retira ...') manda; si no, palabra clave del cliente; si no, el cliente tal cual.
    Devuelve (esperado, esperado_solo_por_cliente)."""
    R = reglas or REGLAS
    o = _sin_acentos(str(obs or "").lower())
    c = _sin_acentos(str(cliente or "").lower())
    por_cliente = None
    for clave, origen in R["origen_por_cliente"]:
        if clave in c:
            por_cliente = origen
            break
    if por_cliente is None:
        por_cliente = str(cliente or "").strip()
    for clave, origen in R["retira"]:
        if clave in o:
            return origen, por_cliente
    return por_cliente, por_cliente


def destino_esperado(obs):
    """3.5: lo que sigue a 'p/'; si no hay, las primeras tres palabras de Observaciones."""
    s = str(obs or "").strip()
    m = re.search(r"p/\s*(.+)", s, re.I)
    if m:
        return m.group(1).strip()
    return " ".join(s.split()[:3])


def _kg(v):
    if v is None:
        return "—"
    s = f"{v:,.0f}".replace(",", ".")
    return s + " kg"


# ───────────────────────── lectura ─────────────────────────
def _color(rgb):
    if not isinstance(rgb, str) or len(rgb) < 6:
        return "otro"
    h = rgb[-6:]
    try:
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    except ValueError:
        return "otro"
    if r >= 245 and g >= 245 and b >= 245:
        return None
    if r >= 200 and g >= 200 and b <= 130:
        return "amarillo"
    if g >= 120 and r <= 130 and b <= 150:
        return "verde"
    if r >= 170 and g <= 100 and b <= 100:
        return "rojo"
    return "otro"


def color_fila(ws, fila, hasta=11):
    """El color no siempre está en la columna A: primer relleno sólido entre las columnas 1 y 11."""
    for c in range(1, hasta + 1):
        f = ws.cell(fila, c).fill
        if f is None or f.fill_type != "solid":
            continue
        fg = f.fgColor
        if fg is None:
            return "otro"
        if getattr(fg, "type", "") == "rgb":
            return _color(fg.rgb)
        if getattr(fg, "type", "") == "theme":
            # un tema claro con tinte 0 suele ser blanco; cualquier otro se trata como "otro" (no se procesa)
            return "otro"
        return "otro"
    return None


def leer_remitos(ruta, avisos=None):
    """Hoja «Remitos»: encabezados en la fila 6, datos desde la 7. Se abre dos veces: estilos y valores calculados."""
    import openpyxl
    avisos = avisos if avisos is not None else []
    wb_e = openpyxl.load_workbook(ruta, data_only=False)
    wb_v = openpyxl.load_workbook(ruta, data_only=True)
    if "Remitos" in wb_e.sheetnames:
        ws_e, ws_v = wb_e["Remitos"], wb_v["Remitos"]
    else:
        ws_e, ws_v = wb_e.worksheets[0], wb_v.worksheets[0]
        avisos.append(f"El archivo de remitos no tiene una hoja llamada «Remitos»; se usó la primera hoja («{ws_e.title}»).")
    # fila de encabezados: la 6, salvo que «Remito» esté en otra de las primeras 12
    cab = 6
    for r in range(1, 13):
        vals = [_txt(ws_v.cell(r, c).value).lower() for c in range(1, 12)]
        if any("remito" in v for v in vals) and any("carga" in v or "cliente" in v for v in vals):
            cab = r
            break
    if cab != 6:
        avisos.append(f"Los encabezados de remitos están en la fila {cab}, no en la 6. Se leyó desde ahí.")
    out, sin_valor = [], 0
    for r in range(cab + 1, ws_v.max_row + 1):
        v = {c: ws_v.cell(r, c).value for c in range(1, 12)}
        remito = _txt(v[7])
        if not remito and not _txt(v[5]) and not _txt(v[8]):
            continue
        carga = _num(v[10])
        if carga is None and ws_e.cell(r, 10).value not in (None, ""):
            sin_valor += 1
        iso, ftxt = _fecha(v[2])
        out.append({"fila": r, "fecha": iso, "fecha_txt": ftxt, "cliente": _txt(v[5]), "obs": _txt(v[6]), "remito": remito,
                    "producto": _txt(v[8]), "carga": carga, "descripcion": _txt(v[11]), "color": color_fila(ws_e, r)})
    if sin_valor:
        avisos.append(f"{sin_valor} fila(s) de remitos tienen la Carga como fórmula sin valor calculado: abrí el archivo en Excel, "
                      f"guardalo y volvé a subirlo. Esas filas quedan «para revisar» sin control de cantidad.")
    if not out:
        avisos.append("El archivo de remitos no tiene filas con datos debajo de los encabezados.")
    return out


def leer_declaraciones(ruta, avisos=None):
    """Export de RENPRE (.xls viejo, también .xlsx o .csv). Detecta la fila de encabezados; sin ella, usa las 7 columnas por posición."""
    import pandas as pd
    avisos = avisos if avisos is not None else []
    ext = ruta.lower().rsplit(".", 1)[-1]
    with open(ruta, "rb") as f:
        inicio = f.read(8)
    if ext == "csv" or (ext not in ("xls", "xlsx", "xlsm") and inicio[:2] not in (b"\xd0\xcf", b"PK")):
        df = pd.read_csv(ruta, header=None, sep=None, engine="python", dtype=str)
    elif inicio[:1] == b"<":
        raise ValueError("el archivo de declaraciones es una página web guardada como .xls; abrilo en Excel y guardalo como .xls o .xlsx")
    elif inicio[:2] == b"PK":
        df = pd.read_excel(ruta, header=None, engine="openpyxl")
    else:
        df = pd.read_excel(ruta, header=None, engine="xlrd")
    cab = None
    for i in range(min(len(df), 40)):
        celdas = [_txt(x).upper() for x in df.iloc[i].tolist()]
        if any("FECHA" in c for c in celdas) and any("CANTIDAD" in c or "PRODUCTO" in c for c in celdas):
            cab = i
            break
    campos = {"id": 0, "fecha": 1, "origen": 2, "destino": 3, "producto": 4, "cantidad": 5, "doc": 6}
    if cab is not None:
        celdas = [_txt(x).upper() for x in df.iloc[cab].tolist()]
        enc = {}
        for j, c in enumerate(celdas):
            if not c:
                continue
            if "ORIGEN" in c:
                enc.setdefault("origen", j)
            elif "DESTINO" in c:
                enc.setdefault("destino", j)
            elif "PRODUCTO" in c:
                enc.setdefault("producto", j)
            elif "CANTIDAD" in c:
                enc.setdefault("cantidad", j)
            elif "DOC" in c or "NUMERO" in c or "NÚMERO" in c:
                enc.setdefault("doc", j)
            elif "FECHA" in c:
                enc.setdefault("fecha", j)
            elif c == "ID":
                enc.setdefault("id", j)
        if len(enc) >= 5:
            campos.update(enc)
        datos = df.iloc[cab + 1:]
        if cab > 0:
            avisos.append(f"Las declaraciones tienen {cab} fila(s) antes de los encabezados; se ignoraron.")
    else:
        datos = df
        avisos.append("No se encontró la fila de encabezados en las declaraciones; se leyeron las siete columnas por posición.")
    out = []
    for _, fila in datos.iterrows():
        g = lambda k: fila.iloc[campos[k]] if campos[k] < len(fila) else None
        idv = _txt(g("id"))
        if not idv and not _txt(g("origen")) and not _txt(g("producto")):
            continue
        iso, ftxt = _fecha(g("fecha"))
        o_raw, d_raw = _txt(g("origen")), _txt(g("destino"))
        o_cuit, o_nom = (o_raw.split(" - ", 1) + [""])[:2] if " - " in o_raw else ("", o_raw)
        d_cuit, d_nom = (d_raw.split(" - ", 1) + [""])[:2] if " - " in d_raw else ("", d_raw)
        prod_raw = _txt(g("producto"))
        prod = prod_raw.split(" - ", 1)[0].strip() if " - " in prod_raw else prod_raw
        out.append({"id": idv, "fecha": iso, "fecha_txt": ftxt, "origen": o_nom.strip(), "origen_cuit": o_cuit.strip(),
                    "destino": d_nom.strip(), "destino_cuit": d_cuit.strip(), "producto": prod,
                    "cantidad": _num(g("cantidad")), "doc": _txt(g("doc")).replace(".0", "") if re.match(r"^\d+\.0$", _txt(g("doc"))) else _txt(g("doc"))})
    if not out:
        avisos.append("El archivo de declaraciones no tiene filas con datos.")
    return out


# ───────────────────────── cruce ─────────────────────────
def _distancia1(a, b):
    """True si a y b difieren en un solo dígito, uno de más/menos o una transposición (DOC mal tipeado)."""
    if a == b:
        return False
    if len(a) == len(b):
        d = [i for i in range(len(a)) if a[i] != b[i]]
        return len(d) == 1 or (len(d) == 2 and d[1] == d[0] + 1 and a[d[0]] == b[d[1]] and a[d[1]] == b[d[0]])
    if abs(len(a) - len(b)) == 1:
        s, l = (a, b) if len(a) < len(b) else (b, a)
        for i in range(len(l)):
            if l[:i] + l[i + 1:] == s:
                return True
    return False


def cruzar(remitos, decls, reglas=None):
    R = reglas or REGLAS
    tol_kg, tol_pct = float(R["tolerancia_kg"]), float(R["tolerancia_pct"]) / 100.0

    validas = [d for d in decls if clave_doc(d["doc"])]
    anuladas = len(decls) - len(validas)
    por_doc = defaultdict(list)
    for d in validas:
        por_doc[clave_doc(d["doc"])].append(d)
    usadas = set()
    # remitos ya controlados (verdes y rojos), para distinguir una declaración huérfana de la de un remito ya visto
    controlados = {}
    for rm in remitos:
        if rm["color"] in ("verde", "rojo"):
            for p in partes_remito(rm["remito"]):
                controlados.setdefault(clave_doc(p), rm)

    amarillas = [rm for rm in remitos if rm["color"] == "amarillo"]
    filas, revisar = [], []
    for rm in amarillas:
        partes = partes_remito(rm["remito"])
        claves = [clave_doc(p) for p in partes]
        combinado = len(partes) > 1
        problemas, notas, encontradas = [], [], []
        faltan = []
        for p, k in zip(partes, claves):
            ds = por_doc.get(k, [])
            if ds:
                encontradas.extend(ds)
            else:
                faltan.append(p)
        base = {"fila": rm["fila"], "fecha": rm["fecha"], "fecha_txt": rm["fecha_txt"], "remito": rm["remito"], "cliente": rm["cliente"],
                "producto": rm["producto"], "carga": rm["carga"], "obs": rm["obs"], "combinado": combinado}
        if not encontradas:
            filas.append(dict(base, estado="PENDIENTE", problemas=[], acciones=["Sin declarar en el mes: declarar."], decls=[]))
            continue
        for d in encontradas:
            usadas.add(id(d))
        if combinado and faltan:
            problemas.append(("pata", f"Falta la pata {', '.join(faltan)} del combinado: declararla."))

        # 3.3 origen y 3.5 destino, por declaración
        esp_o, esp_o_cli = origen_esperado(rm["cliente"], rm["obs"], R)
        t_eo, t_eo_cli = tokens(esp_o, R), tokens(esp_o_cli, R)
        esp_d = destino_esperado(rm["obs"])
        t_ed = tokens(esp_d, R)
        fam_rem = familia(rm["producto"], R)
        if not t_ed:
            notas.append("Sin destino en Observaciones (falta «p/ …»): no se pudo controlar el destino.")
        for d in encontradas:
            t_o, t_d = tokens(d["origen"], R), tokens(d["destino"], R)
            ok_o = bool(t_eo & t_o) or (combinado and bool(t_eo_cli & t_o))
            ok_d = (not t_ed) or bool(t_ed & t_d)
            etq = f" (doc {d['doc']})" if combinado else ""
            if not ok_o and not ok_d and t_ed and (t_eo & t_d) and (t_ed & t_o):
                problemas.append(("invertido", f"Origen y destino invertidos{etq} (declarado {d['origen']} → {d['destino']}): anular y re-cargar."))
            else:
                if not ok_o:
                    problemas.append(("origen", f"Origen declarado «{d['origen']}»{etq}, esperado «{esp_o}». Anular y re-cargar con el origen correcto."))
                if not ok_d:
                    problemas.append(("destino", f"Destino declarado «{d['destino']}»{etq}, esperado «{esp_d}». Anular y re-cargar con el destino correcto."))
            # 3.4 producto por familia
            fam_d = familia(d["producto"], R)
            if fam_d != fam_rem and fam_rem not in R["familias_libres"]:
                problemas.append(("producto", f"Producto declarado «{d['producto']}»{etq} y el remito es «{rm['producto']}». Anular y re-cargar con el producto correcto."))
            # 4: alerta de fecha
            if d["fecha"]:
                try:
                    dia = datetime.date.fromisoformat(d["fecha"]).weekday()
                    if dia == 6:
                        notas.append(f"La declaración {d['id']} está fechada un domingo ({d['fecha_txt']}): revisar la fecha.")
                except ValueError:
                    pass

        # 3.6 cantidad y 3.7 duplicados
        suma = sum(d["cantidad"] or 0 for d in encontradas)
        n_decl = len(encontradas)
        if not combinado:
            if rm["carga"] is None:
                notas.append("La Carga del remito no tiene valor (fórmula sin calcular): no se controló la cantidad.")
            else:
                tol = max(tol_kg, tol_pct * rm["carga"])
                if n_decl > 1:
                    fams = {familia(d["producto"], R) for d in encontradas}
                    dests = [tokens(d["destino"], R) for d in encontradas]
                    distintos = len(fams) > 1 or any(not (dests[0] & x) for x in dests[1:])
                    if distintos:
                        problemas.append(("duplicado", f"DUPLICADO: doc {rm['remito']} declarado {n_decl} veces con distinto producto o destino. Anular la línea incorrecta."))
                    elif suma - rm["carga"] > tol:
                        problemas.append(("duplicado", f"DUPLICADO: doc {rm['remito']} declarado {n_decl} veces; las líneas suman {_kg(suma)} contra una carga de {_kg(rm['carga'])}. Anular la línea sobrante."))
                    elif abs(suma - rm["carga"]) <= tol:
                        notas.append(f"Declaración parcial en {n_decl} líneas que suman {_kg(suma)} (carga {_kg(rm['carga'])}): válida.")
                    else:
                        problemas.append(("cantidad", f"Las {n_decl} líneas suman {_kg(suma)} contra una carga de {_kg(rm['carga'])}. Corregir la cantidad."))
                elif abs(suma - rm["carga"]) > tol:
                    problemas.append(("cantidad", f"Cantidad declarada {_kg(suma)} contra una carga de {_kg(rm['carga'])}. Corregir la cantidad."))
        else:
            docs_rep = [k for k in claves if len(por_doc.get(k, [])) > 1]
            if docs_rep:
                notas.append(f"Combinado con doc(s) {', '.join(docs_rep)} declarados más de una vez: revisar que no haya línea de más.")
            if rm["carga"] is not None and n_decl and abs(suma - rm["carga"]) > max(tol_kg, tol_pct * rm["carga"]):
                notas.append(f"Combinado: las declaraciones suman {_kg(suma)} y la carga es {_kg(rm['carga'])} (no se exige que cierre).")

        estado = "PROBLEMA" if problemas else "OK"
        filas.append(dict(base, estado=estado, problemas=[p[0] for p in problemas], acciones=[p[1] for p in problemas], notas=notas,
                          decls=[{"id": d["id"], "fecha_txt": d["fecha_txt"], "doc": d["doc"], "origen": d["origen"], "destino": d["destino"],
                                  "producto": d["producto"], "cantidad": d["cantidad"]} for d in encontradas]))
        if notas:
            revisar.append(dict(base, estado=estado, motivos=notas))

    # 4: declaraciones que ninguna fila amarilla usó
    pendientes_claves = {clave_doc(p): f["remito"] for f in filas if f["estado"] == "PENDIENTE" for p in partes_remito(f["remito"])}
    cargas = {}
    for rm in amarillas:
        if rm["carga"] is not None:
            cargas.setdefault(str(int(round(rm["carga"]))), rm["remito"])
    sin_remito = []
    for d in validas:
        if id(d) in usadas:
            continue
        k = clave_doc(d["doc"])
        if k in controlados:
            rm = controlados[k]
            causa = f"Es de un remito ya controlado (fila {rm['fila']}, {'verde' if rm['color'] == 'verde' else 'roja'}): no hace falta nada."
            tipo = "controlado"
        else:
            pistas = []
            parecidos = [r for kp, r in pendientes_claves.items() if _distancia1(k, kp)]
            if parecidos:
                pistas.append(f"¿será el remito {', '.join(sorted(set(parecidos)))} (sin declarar)? Puede estar mal tipeado el número.")
            if k in cargas:
                pistas.append(f"el número coincide con la Carga del remito {cargas[k]}: puede haberse cargado la cantidad en el campo del remito.")
            causa = " ".join(p[0].upper() + p[1:] for p in pistas) if pistas else "No corresponde a ningún remito del archivo: puede ser de otro mes, un número mal tipeado o un duplicado anterior."
            tipo = "huerfana"
        sin_remito.append({"id": d["id"], "fecha_txt": d["fecha_txt"], "fecha": d["fecha"], "doc": d["doc"], "origen": d["origen"], "destino": d["destino"],
                           "producto": d["producto"], "cantidad": d["cantidad"], "causa": causa, "tipo": tipo})

    mal = [f for f in filas if f["estado"] != "OK"]
    bien = [f for f in filas if f["estado"] == "OK"]
    res = {"amarillas": len(amarillas), "ok": len(bien), "problema": sum(1 for f in filas if f["estado"] == "PROBLEMA"),
           "pendiente": sum(1 for f in filas if f["estado"] == "PENDIENTE"), "revisar": len(revisar),
           "decl_total": len(decls), "decl_validas": len(validas), "anuladas": anuladas,
           "sin_remito": sum(1 for s in sin_remito if s["tipo"] == "huerfana"), "sin_remito_controladas": sum(1 for s in sin_remito if s["tipo"] == "controlado"),
           "verdes": sum(1 for rm in remitos if rm["color"] == "verde"), "rojas": sum(1 for rm in remitos if rm["color"] == "rojo"),
           "otras": sum(1 for rm in remitos if rm["color"] not in ("amarillo", "verde", "rojo"))}
    return {"resumen": res, "mal": mal, "bien": bien, "revisar": revisar, "sin_remito": sin_remito}


def mes_de(remitos, decls):
    cuenta = defaultdict(int)
    for x in list(remitos) + list(decls):
        if x.get("fecha"):
            cuenta[x["fecha"][:7]] += 1
    if not cuenta:
        return ""
    ym = max(cuenta, key=cuenta.get)
    return f"{MESES[int(ym[5:7])]} {ym[:4]}"


# ───────────────────────── pantalla ─────────────────────────
def _e(s):
    return html.escape("" if s is None else str(s))


def _kpi(valor, rotulo, nota=""):
    return f'<div class="kpi"><div class="kv">{valor}</div><div class="kl">{rotulo}</div>{f"<div class=kn>{nota}</div>" if nota else ""}</div>'


def _tabla(cols, filas):
    th = "".join(f'<th class="{a}">{t}</th>' for t, a in cols)
    trs = "".join("<tr>" + "".join(f'<td class="{cols[i][1]}">{c}</td>' for i, c in enumerate(f)) + "</tr>" for f in filas)
    return f'<div class="tw"><table><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table></div>'


def _badge(estado):
    col = {"PROBLEMA": "#C62828", "PENDIENTE": "#E65100", "OK": "#2E7D32"}[estado]
    etq = {"PROBLEMA": "Con problema", "PENDIENTE": "Sin declarar", "OK": "Bien"}[estado]
    return f'<span class="est" style="background:{col}">{etq}</span>'


def _ap(pid, titulo, bajada, cuerpo):
    return f'<article class="apartado" id="{pid}" data-ap="{pid}" hidden><header class="aph"><h1>{titulo}</h1><p>{bajada}</p></header>{cuerpo}</article>'


def render_html(R, mes, nombres, avisos):
    r = R["resumen"]
    tabs = [("rp-resumen", "Resumen"), ("rp-mal", f"Con problemas · {r['problema'] + r['pendiente']}"), ("rp-revisar", f"Para revisar · {r['revisar']}"),
            ("rp-bien", f"Bien declarados · {r['ok']}"), ("rp-sin", f"Sin remito · {r['sin_remito']}")]
    t = "".join(f'<button type="button" data-ir="{i}">{_e(n)}</button>' for i, n in tabs)
    partes = []

    # Resumen
    av = "".join(f'<p class="aviso">{_e(a)}</p>' for a in avisos)
    kp = (_kpi(r["amarillas"], "Remitos a controlar", "filas amarillas") + _kpi(r["ok"], "Bien declarados") + _kpi(r["problema"], "Con problema")
          + _kpi(r["pendiente"], "Sin declarar") + _kpi(r["revisar"], "Para revisar", "pasan, pero conviene mirarlos") + _kpi(r["sin_remito"], "Declaraciones sin remito"))
    hechos = [f"Se leyeron {r['decl_total']} declaraciones: {r['decl_validas']} con número de documento y {r['anuladas']} anuladas (sin número), que no se cruzan.",
              f"El archivo de remitos trae {r['verdes']} fila(s) verde(s) y {r['rojas']} roja(s) ya controladas, que no se vuelven a procesar"
              + (f", y {r['otras']} sin color o con otro color, que tampoco." if r["otras"] else ".")]
    if r["sin_remito_controladas"]:
        hechos.append(f"{r['sin_remito_controladas']} declaración(es) corresponden a remitos ya controlados (verdes o rojos): figuran al final de «Sin remito» solo como referencia.")
    lead = (f"De los {r['amarillas']} remitos pendientes de control, {r['ok']} están bien declarados, {r['problema']} tienen algún problema y {r['pendiente']} no fueron declarados."
            if r["amarillas"] else "El archivo de remitos no tiene filas amarillas: no hay nada pendiente de control.")
    partes.append(_ap("rp-resumen", f"Control RENPRE · {mes or 'mes'}", f"Remitos: {_e(nombres[0])} · Declaraciones: {_e(nombres[1])}. Reglas v{VERSION}.",
                      f'<section id="rp-resumen-s"><h2>Resumen</h2>{av}<p class="lead">{lead}</p><div class="kpis">{kp}</div><ul class="hechos">'
                      + "".join(f"<li>{_e(h)}</li>" for h in hechos) + "</ul></section>"))

    # Con problemas
    cols = [("Fecha", "l"), ("Remito", "l"), ("Cliente", "l"), ("Producto", "l"), ("Carga", "r"), ("Estado", "l"), ("Qué hacer", "l")]
    filas = []
    for f in R["mal"]:
        acc = "<br>".join(_e(a) for a in f["acciones"])
        filas.append([_e(f["fecha_txt"]), f"<b>{_e(f['remito'])}</b>", _e(f["cliente"]), _e(f["producto"]), _kg(f["carga"]), _badge(f["estado"]), f'<span class="obs">{acc}</span>'])
    cuerpo = (_tabla(cols, filas) if filas else '<p class="nota">No hay remitos con problemas ni sin declarar.</p>')
    partes.append(_ap("rp-mal", "Con problemas", "Remitos con alguna diferencia contra lo declarado, y los que no fueron declarados. Cada uno dice qué hacer.",
                      f'<section id="rp-mal-s"><h2>Remitos a corregir</h2>{cuerpo}</section>'))

    # Para revisar
    cols = [("Fecha", "l"), ("Remito", "l"), ("Cliente", "l"), ("Carga", "r"), ("Estado", "l"), ("Motivo", "l")]
    filas = [[_e(f["fecha_txt"]), f"<b>{_e(f['remito'])}</b>", _e(f["cliente"]), _kg(f["carga"]), _badge(f["estado"]), f'<span class="obs">{"<br>".join(_e(m) for m in f["motivos"])}</span>'] for f in R["revisar"]]
    cuerpo = (_tabla(cols, filas) if filas else '<p class="nota">Nada para revisar aparte de lo que está en «Con problemas».</p>')
    partes.append(_ap("rp-revisar", "Para revisar", "Casos que pasan el control pero conviene mirar: combinados con diferencia de cantidad, parciales, fechas raras.",
                      f'<section id="rp-revisar-s"><h2>Casos borderline</h2>{cuerpo}</section>'))

    # Bien
    cols = [("Fecha", "l"), ("Remito", "l"), ("Cliente", "l"), ("Producto", "l"), ("Carga", "r"), ("Destino declarado", "l"), ("Cantidad declarada", "r")]
    filas = []
    for f in R["bien"]:
        dest = " / ".join(sorted({d["destino"] for d in f["decls"]}))
        cant = sum(d["cantidad"] or 0 for d in f["decls"])
        filas.append([_e(f["fecha_txt"]), f"<b>{_e(f['remito'])}</b>", _e(f["cliente"]), _e(f["producto"]), _kg(f["carga"]), _e(dest), _kg(cant)])
    cuerpo = (_tabla(cols, filas) if filas else '<p class="nota">Ningún remito pasó el control completo.</p>')
    partes.append(_ap("rp-bien", "Bien declarados", "Remitos cuya declaración coincide en origen, destino, producto y cantidad.",
                      f'<section id="rp-bien-s"><h2>Remitos bien declarados</h2>{cuerpo}</section>'))

    # Sin remito
    cols = [("ID", "l"), ("Fecha", "l"), ("Doc", "l"), ("Origen", "l"), ("Destino", "l"), ("Producto", "l"), ("Cantidad", "r"), ("Posible causa", "l")]
    orden = sorted(R["sin_remito"], key=lambda s: (s["tipo"] == "controlado", s["fecha"]))
    filas = [[_e(s["id"]), _e(s["fecha_txt"]), f"<b>{_e(s['doc'])}</b>", _e(s["origen"]), _e(s["destino"]), _e(s["producto"]), _kg(s["cantidad"]), f'<span class="obs">{_e(s["causa"])}</span>'] for s in orden]
    cuerpo = (_tabla(cols, filas) if filas else '<p class="nota">Todas las declaraciones con número de documento corresponden a un remito del archivo.</p>')
    partes.append(_ap("rp-sin", "Declaraciones sin remito", "Declaraciones con número de documento que ningún remito pendiente usó: cruce de mes, número mal tipeado o duplicado anterior.",
                      f'<section id="rp-sin-s"><h2>Declaraciones que no cruzan</h2>{cuerpo}</section>'))

    return f'<template data-tabs>{t}</template><div data-inicio="rp-resumen"></div>' + "".join(partes)


def texto(R, mes):
    """Salida de texto (sección 5), para el control desde la consola."""
    r = R["resumen"]
    L = [f"CONTROL RENPRE · {mes}", f"Amarillas: {r['amarillas']} · OK: {r['ok']} · con problema: {r['problema']} · sin declarar: {r['pendiente']} · para revisar: {r['revisar']}",
         f"Declaraciones: {r['decl_total']} ({r['decl_validas']} válidas, {r['anuladas']} anuladas) · sin remito: {r['sin_remito']}", "", "MAL"]
    for f in R["mal"]:
        L.append(f"  {f['fecha_txt']} · {f['remito']} · {f['cliente']} · {_kg(f['carga'])}")
        for a in f["acciones"]:
            L.append(f"      - {a}")
    L += ["", "BIEN"]
    for f in R["bien"]:
        L.append(f"  {f['fecha_txt']} · {f['remito']} · {f['cliente']} · {f['producto']} · {_kg(f['carga'])} · → {' / '.join(sorted({d['destino'] for d in f['decls']}))}")
    L += ["", "SIN REMITO"]
    for s in R["sin_remito"]:
        L.append(f"  {s['id']} · {s['fecha_txt']} · doc {s['doc']} · {s['origen']} → {s['destino']} · {s['producto']} · {_kg(s['cantidad'])} · {s['causa']}")
    L += ["", "PARA REVISAR"]
    for f in R["revisar"]:
        L.append(f"  {f['fecha_txt']} · {f['remito']} · {f['cliente']}: " + " | ".join(f["motivos"]))
    return "\n".join(L)


def procesar(ruta_remitos, ruta_decl, nombres=None, reglas=None):
    """Punto de entrada único (consola y navegador). Devuelve dict con resumen, listas, avisos, errores y html."""
    nombres = nombres or (ruta_remitos.rsplit("/", 1)[-1], ruta_decl.rsplit("/", 1)[-1])
    avisos, errores = [], []
    remitos, decls = [], []
    try:
        remitos = leer_remitos(ruta_remitos, avisos)
    except Exception as ex:
        errores.append(f"No se pudo leer el archivo de remitos ({nombres[0]}): {type(ex).__name__}: {ex}")
    try:
        decls = leer_declaraciones(ruta_decl, avisos)
    except Exception as ex:
        errores.append(f"No se pudo leer el archivo de declaraciones ({nombres[1]}): {type(ex).__name__}: {ex}")
    if errores:
        return {"errores": errores, "avisos": avisos, "html": "", "resumen": None}
    R = cruzar(remitos, decls, reglas)
    mes = mes_de(remitos, decls)
    R["mes"] = mes
    R["avisos"] = avisos
    R["errores"] = []
    R["html"] = render_html(R, mes, nombres, avisos)
    R["texto"] = texto(R, mes)
    return R


def procesar_json(ruta_remitos, ruta_decl, nombre_remitos="", nombre_decl=""):
    """Para el navegador: mismo resultado, como JSON."""
    nombres = (nombre_remitos or ruta_remitos.rsplit("/", 1)[-1], nombre_decl or ruta_decl.rsplit("/", 1)[-1])
    return json.dumps(procesar(ruta_remitos, ruta_decl, nombres), ensure_ascii=False, default=str)
